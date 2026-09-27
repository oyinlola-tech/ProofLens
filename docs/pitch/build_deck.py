"""Build the ProofLens pitch deck (.pptx) from real screenshots, for import into Google Slides."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SP = Path(__file__).parent
SHOTS = SP / "screenshots"
CROPS = Path("/tmp") / "prooflens-deck-crops"
CROPS.mkdir(exist_ok=True)

# ProofLens brand, taken from apps/web/src/app/globals.css and docs/brand.
BG = RGBColor(0x11, 0x11, 0x13)
SURFACE = RGBColor(0x1B, 0x1B, 0x1F)
SURFACE_2 = RGBColor(0x24, 0x24, 0x29)
LINE = RGBColor(0x33, 0x33, 0x3A)
TEXT = RGBColor(0xF2, 0xF0, 0xEA)
MUTED = RGBColor(0x9A, 0x9A, 0xA3)
ACCENT = RGBColor(0xFF, 0x6A, 0x3D)
ACCENT_DEEP = RGBColor(0xE5, 0x50, 0x2A)
RED = RGBColor(0xF0, 0x5D, 0x5E)
GREEN = RGBColor(0x3D, 0xD6, 0x8C)
AMBER = RGBColor(0xF5, 0xB8, 0x4A)
BLUE = RGBColor(0x7A, 0xA2, 0xFF)

HEAD = "Bricolage Grotesque"
BODY = "Geist"
MONO = "Geist Mono"


def crop(name: str, box: tuple[int, int, int, int], out: str) -> str:
    img = Image.open(SHOTS / name)
    path = CROPS / out
    img.crop(box).save(path, quality=95)
    return str(path)


# Screenshots are 1536x692. The left 256px is the app sidebar (and the Next.js dev badge).
C = {
    "landing": crop("01-landing.jpg", (0, 0, 1520, 692), "landing.jpg"),
    "verdict": crop("09-verdict.jpg", (330, 30, 1440, 692), "verdict.jpg"),
    "why": crop("10-why-show-me-why.jpg", (330, 10, 1440, 692), "why.jpg"),
    "checks": crop("11-conclude-checks.jpg", (330, 10, 1440, 680), "checks.jpg"),
    "viewer": crop("12-source-viewer.jpg", (340, 30, 1440, 420), "viewer.jpg"),
    "otp": crop("03-otp.jpg", (500, 80, 1020, 530), "otp.jpg"),
    "evidence": crop("07-evidence-attached.jpg", (330, 30, 1440, 590), "evidence.jpg"),
    "picker": crop("06-passage-picker.jpg", (350, 0, 1070, 650), "picker.jpg"),
    "history": crop("13-history.jpg", (330, 30, 1450, 240), "history.jpg"),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def box(slide, x, y, w, h, fill=SURFACE, line=None, radius=True, weight=1):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if radius:
        shp.adjustments[0] = 0.06
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(weight)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, runs, size=18, color=TEXT, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.12):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = para if isinstance(para, list) else [(para, {})]
        for part in parts:
            t, o = (part, {}) if isinstance(part, str) else part
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
        if isinstance(para, list) and para and isinstance(para[0], tuple) and "space_after" in para[0][1]:
            p.space_after = Pt(para[0][1]["space_after"])
    return tb


def pill(slide, x, y, label, fg, bgc, w=None, size=11):
    w = w or (0.0075 * size * len(label) + 0.4)
    b = box(slide, x, y, w, 0.34, fill=bgc)
    b.adjustments[0] = 0.5
    text(slide, x, y, w, 0.34, label, size=size, color=fg, font=BODY, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def mark(slide, x, y, d=0.42):
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    ring.fill.background()
    ring.line.color.rgb = TEXT
    ring.line.width = Pt(d * 9)
    dot_d = d * 0.5
    dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + (d - dot_d) / 2), Inches(y + (d - dot_d) / 2), Inches(dot_d), Inches(dot_d)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = ACCENT_DEEP
    dot.line.fill.background()


def header(slide, n, kicker, title, sub=None):
    mark(slide, 0.6, 0.5, 0.28)
    text(slide, 0.98, 0.5, 4, 0.3, "ProofLens", size=13, color=TEXT, font=HEAD, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 11.4, 0.5, 1.33, 0.3, f"{n:02d} / 10", size=11, color=MUTED, font=MONO,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.6, 1.2, 12, 0.3, kicker.upper(), size=12, color=ACCENT, font=MONO, bold=True)
    text(slide, 0.6, 1.55, 12.1, 1.0, title, size=36, color=TEXT, font=HEAD, bold=True, spacing=1.0)
    if sub:
        text(slide, 0.6, 2.45, 11.5, 0.6, sub, size=16, color=MUTED)


def picture(slide, path, x, y, w=None, h=None, frame=True):
    if frame:
        img = Image.open(path)
        ratio = img.height / img.width
        ww = w if w else h / ratio
        hh = h if h else w * ratio
        box(slide, x - 0.06, y - 0.06, ww + 0.12, hh + 0.12, fill=SURFACE_2, line=LINE)
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def caption(slide, x, y, w, label):
    text(slide, x, y, w, 0.3, label, size=11, color=MUTED, font=MONO)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def arrow(slide, x, y, w=0.5, color=ACCENT):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.3))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()


def down_arrow(slide, x, y, color=ACCENT):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(0.3), Inches(0.35))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
bg(s)
glow = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.2), Inches(-2.5), Inches(8.5), Inches(8.5))
glow.fill.solid()
glow.fill.fore_color.rgb = RGBColor(0x2A, 0x17, 0x11)
glow.line.fill.background()
mark(s, 0.8, 0.8, 0.6)
text(s, 1.55, 0.8, 5, 0.6, "ProofLens", size=24, font=HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 2.0, 6.6, 2.4, [
    [("Check what your ", {}), ("evidence", {"color": ACCENT}), (" actually supports.", {})],
], size=44, font=HEAD, bold=True, spacing=0.98)
text(s, 0.8, 4.75, 6.0, 1.2,
     "Evidence-grounded claim verification. Enter a claim, add your documents, and get a verdict "
     "that explains itself and points back to the page it came from.",
     size=17, color=MUTED, spacing=1.25)
text(s, 0.8, 6.45, 7, 0.3, "GOMYCODE NIGERIA  ·  COME BUILD WITH AI HACKATHON 2026",
     size=11, color=MUTED, font=MONO, bold=True)
picture(s, C["verdict"], 7.55, 1.55, w=5.25)
caption(s, 7.55, 4.85, 5.3, "Real result screen · live run, 27 Sep 2026")
notes(s, """
ProofLens answers one narrow question: does the evidence I supplied actually support the claim I am making?
You give it a claim and your source documents. It gives back a verdict, an explanation of what the source does and does not establish, and the exact passage and page it relied on.
Everything on these slides comes from the running product and the repository. The screenshot on the right is a real result from a live run today.
""")

# ---------------------------------------------------------------- 2. Problem
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 2, "The problem", "A source can be real and still not say what the claim says.")
box(s, 0.6, 2.85, 5.6, 2.6, fill=SURFACE, line=LINE)
text(s, 0.95, 3.1, 5, 0.3, "THE CLAIM", size=11, color=MUTED, font=MONO, bold=True)
text(s, 0.95, 3.5, 5.0, 1.8, [
    [("“The study ", {}), ("proves", {"color": ACCENT, "bold": True}), (" that drinking coffee every day ", {}),
     ("causes", {"color": ACCENT, "bold": True}), (" a lower risk of heart disease.”", {})],
], size=22, font=HEAD, spacing=1.15)
box(s, 7.1, 2.85, 5.6, 2.6, fill=SURFACE, line=LINE)
text(s, 7.45, 3.1, 5, 0.3, "THE SOURCE  ·  PAGE 2 AND PAGE 3", size=11, color=MUTED, font=MONO, bold=True)
text(s, 7.45, 3.5, 5.0, 1.8, [
    [("“Daily coffee consumption was ", {}), ("associated with", {"color": GREEN, "bold": True}),
     (" a lower risk of heart disease.”", {})],
    [("“…the findings ", {}), ("do not establish", {"color": RED, "bold": True}), (" that coffee causes a lower risk.”", {})],
], size=17, spacing=1.2)
arrow(s, 6.4, 4.0, 0.5)
text(s, 0.6, 5.8, 12.1, 0.4, "The gap is easy to miss when reading quickly:", size=15, color=MUTED)
x = 0.6
for label in ["Stronger language than the source", "A different number, date or name", "A negation read as agreement"]:
    w = pill(s, x, 6.3, label, TEXT, SURFACE_2, size=12)
    x += w + 0.2
notes(s, """
The problem is not fake sources. It is real sources being stretched.
A study finds an association, and the claim says it proves causation. A report says 1,200 participants and the claim says 10,000. A sentence says something is not established, and it gets quoted as if it were.
The source exists, the citation looks fine, and the claim still is not supported. Checking that by hand means reading closely every time.
The example here is the one we will run live: a sample document written for the demo.
""")

# ---------------------------------------------------------------- 3. Solution
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 3, "The product", "Claim + evidence in. An explained verdict out.",
       "ProofLens tests a claim against the documents you supply and explains what that evidence establishes.")
y = 3.35
cards = [("Your claim", "One plain statement", ACCENT), ("Your evidence", "PDF, TXT, MD or CSV · passage + page", BLUE)]
box(s, 0.6, y, 3.4, 1.15, fill=SURFACE, line=LINE)
text(s, 0.85, y + 0.2, 3, 0.4, "Your claim", size=20, font=HEAD, bold=True)
text(s, 0.85, y + 0.62, 3, 0.4, "One plain statement", size=13, color=MUTED)
text(s, 2.15, y + 1.2, 0.5, 0.5, "+", size=26, color=MUTED, font=HEAD, align=PP_ALIGN.CENTER)
box(s, 0.6, y + 1.75, 3.4, 1.15, fill=SURFACE, line=LINE)
text(s, 0.85, y + 1.95, 3, 0.4, "Your evidence", size=20, font=HEAD, bold=True)
text(s, 0.85, y + 2.37, 3, 0.4, "PDF, TXT, MD or CSV · passage + page", size=13, color=MUTED)
arrow(s, 4.2, y + 1.3, 0.6)
core = box(s, 5.0, y + 0.55, 2.5, 1.8, fill=ACCENT_DEEP)
text(s, 5.0, y + 0.55, 2.5, 1.8, [[("ProofLens", {"size": 24, "bold": True, "font": HEAD})],
                                   [("checks + reasoning", {"size": 13, "color": RGBColor(0xFF, 0xE3, 0xD8)})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 7.7, y + 1.3, 0.6)
outs = [("Verdict + confidence", "One of four outcomes", TEXT),
        ("Explanation", "What the source says, why it does or does not match, what it does not establish", TEXT),
        ("Source evidence", "The exact quote, document and page: “Show me why”", TEXT)]
for i, (t, d, c) in enumerate(outs):
    yy = y - 0.05 + i * 1.02
    box(s, 8.5, yy, 4.2, 0.9, fill=SURFACE, line=ACCENT if i == 1 else LINE, weight=1.5 if i == 1 else 1)
    text(s, 8.72, yy + 0.12, 3.9, 0.3, t, size=15, font=HEAD, bold=True)
    text(s, 8.72, yy + 0.44, 3.9, 0.45, d, size=11, color=MUTED, spacing=1.05)
x = 0.6
text(s, 0.6, 6.55, 2.2, 0.34, "THE FOUR VERDICTS", size=11, color=MUTED, font=MONO, bold=True, anchor=MSO_ANCHOR.MIDDLE)
x = 2.75
for label, fg in [("Supported", GREEN), ("Partially supported", AMBER), ("Contradicted", RED), ("Insufficient evidence", MUTED)]:
    w = pill(s, x, 6.55, label, fg, SURFACE_2, size=12)
    x += w + 0.18
notes(s, """
ProofLens is not a chatbot and not a web search. It only reasons over the evidence you attach.
The label is the least interesting part. The explanation is the product: what the source actually says, why the claim does or does not match it, what the source does not establish, and a link to the exact passage.
There are four verdicts. Insufficient evidence is its own outcome, so "the source does not cover this" is never confused with "the source says the opposite".
""")

# ---------------------------------------------------------------- 4. How it works
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 4, "How it works", "Deterministic checks and AI reasoning, with the backend owning the verdict.")
stages = [
    ("01", "Claim analysed", "Strength of language, causal wording, numbers, dates, names"),
    ("02", "Evidence extracted", "PDF text pulled per page in an isolated subprocess; you pick the passage and page"),
    ("03", "Rule-based checks", "Number, date, entity, negation and causation-vs-association comparisons"),
    ("04", "AI reasoning", "Model sees only the supplied passages and must answer in a strict JSON schema"),
    ("05", "Backend validation", "Output parsed with Pydantic; references must match a supplied passage; hard conflicts override the model"),
    ("06", "Show me why", "Verdict with quotes, document and page, linked to the highlighted passage"),
]
cw, gap, top = 1.93, 0.12, 3.05
for i, (n, t, d) in enumerate(stages):
    x = 0.6 + i * (cw + gap)
    hi = i in (2, 3, 4)
    box(s, x, top, cw, 3.35, fill=SURFACE, line=ACCENT if hi else LINE, weight=1.5 if hi else 1)
    text(s, x + 0.2, top + 0.22, cw - 0.4, 0.3, n, size=12, color=ACCENT, font=MONO, bold=True)
    text(s, x + 0.2, top + 0.6, cw - 0.4, 0.8, t, size=15, font=HEAD, bold=True, spacing=1.0)
    text(s, x + 0.2, top + 1.45, cw - 0.4, 1.8, d, size=12, color=MUTED, spacing=1.2)
text(s, 0.6, 6.7, 12.1, 0.4,
     "If every AI model fails, ProofLens returns the deterministic result only when the checks are decisive; otherwise it says reasoning is unavailable and stores nothing.",
     size=12, color=MUTED, font=BODY)
notes(s, """
This pipeline is what the code actually does.
First the claim is analysed: is it absolute or hedged, does it assert causation, what numbers, dates and names does it contain.
Then deterministic checks compare those against the evidence directly, without any model.
The AI model then reasons over the supplied passages only, and has to answer in a strict JSON schema.
The backend validates that answer. References to passages that were not supplied are dropped. Quotes have to appear verbatim in the stored passage. And a hard conflict found by the checks, such as a number mismatch, overrides a model that says supported.
If the AI is unavailable, the user sees an unavailable state. We never show an invented verdict.
""")

# ---------------------------------------------------------------- 5. The important difference
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 5, "The important difference", "Not just a label. What the source establishes, and why.")
cols = [
    ("YOUR CLAIM", "“The study proves that drinking coffee every day causes a lower risk of heart disease.”", ACCENT),
    ("WHAT THE EVIDENCE SAYS", "Daily coffee consumption is associated with an 18% lower incidence of heart disease, but because the study is observational, it does not establish that coffee causes a lower risk.", BLUE),
    ("WHY IT DOES NOT FULLY MATCH", "The source describes an association and notes the observational design prevents establishing causation, whereas the claim asserts proven causal effect.", AMBER),
    ("WHAT YOU CAN CONCLUDE", "Coffee consumption is linked to a lower risk of heart disease, but the study does not prove that coffee causes this reduction.", GREEN),
]
cw = 2.95
for i, (k, v, c) in enumerate(cols):
    x = 0.6 + i * (cw + 0.13)
    box(s, x, 2.9, cw, 2.75, fill=SURFACE, line=LINE)
    bar = box(s, x, 2.9, cw, 0.07, fill=c, radius=False)
    text(s, x + 0.22, 3.15, cw - 0.44, 0.3, k, size=11, color=c, font=MONO, bold=True)
    text(s, x + 0.22, 3.55, cw - 0.44, 2.0, v, size=13.5, color=TEXT, spacing=1.2)
box(s, 0.6, 5.85, 12.1, 1.05, fill=SURFACE_2, line=None)
pill(s, 0.85, 6.2, "Contradicted · 97%", RED, RGBColor(0x3A, 0x1C, 0x1E), w=2.2, size=12)
text(s, 3.3, 5.98, 9.2, 0.8, [
    [("SHOW ME WHY  ·  coffee-heart-study-sample.pdf  ·  page 3  ·  Limitations", {"size": 10, "color": MUTED, "font": MONO, "bold": True})],
    [("“Because this was an observational study, the findings do not establish that coffee causes a lower risk of heart disease.”", {"size": 14})],
], spacing=1.25)
notes(s, """
This is the actual output from the live run, word for word. Model wording varies from run to run, but these fields are always there.
A plain label would just say contradicted. ProofLens explains the gap: the source found an association, the claim says the study proves causation.
It also says what you can honestly conclude from this source, which is the useful part for the person writing the claim.
And Show me why points to the exact sentence on page 3 that the verdict rests on.
The deterministic checks flagged the same thing independently: stronger language than the source, causation versus associated.
The document is a sample written for this demo, and it is labelled as fictional data on page 1.
""")

# ---------------------------------------------------------------- 6. Real product
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 6, "The real product", "The running web app, captured today.")
picture(s, C["why"], 0.6, 2.55, w=6.9)
caption(s, 0.6, 6.9, 7.3, "Result: why it does not match  ·  Show me why, with roles and page numbers")
picture(s, C["viewer"], 8.0, 2.55, w=4.7)
caption(s, 8.0, 4.3, 4.4, "Open the source: the cited sentence highlighted on page 3")
picture(s, C["evidence"], 8.0, 4.9, w=2.6)
picture(s, C["otp"], 10.85, 4.9, h=1.6)
caption(s, 8.0, 6.7, 4.7, "Evidence by page  ·  6-digit email code")
notes(s, """
These are real screenshots from the web app running locally today, not mock-ups.
Large on the left: the result page. Why the claim does not fully match, the unsupported portion, and the Show me why column, where each cited passage has a role, supports or conflicts, and a page number.
Top right: clicking open the source takes you to the document viewer on that page with the cited sentence highlighted.
Bottom right: evidence is attached by page and passage, and accounts are confirmed with a 6-digit email code.
There is also an Expo mobile app with the same flow; we have not captured it on a device, so it is not shown here.
""")

# ---------------------------------------------------------------- 7. Technical foundation
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 7, "Technical foundation", "A full-stack product, not a prompt wrapper.")
cols = [
    ("BACKEND", ["FastAPI · Python", "PostgreSQL · async SQLAlchemy 2", "Alembic migrations (10 revisions)", "PyMuPDF text extraction in a sandboxed subprocess"]),
    ("CLIENTS", ["Next.js 16 · React 19 web app", "Expo SDK 57 mobile app", "Neither client calls an AI provider or holds verification logic"]),
    ("AI LAYER", ["One provider interface: Gemini, NVIDIA NIM, Groq, Ollama", "Chosen by configuration, with ordered fallback models", "Strict JSON-schema output, validated with Pydantic"]),
    ("SECURITY", ["Argon2id passwords · SHA-256 hashed session tokens", "Revocation, 5-session cap, 24h expiry", "Rate limits on login, signup, OTP · owner checks on every resource", "Request body caps · security headers"]),
]
cw = 2.95
for i, (k, items) in enumerate(cols):
    x = 0.6 + i * (cw + 0.13)
    box(s, x, 2.9, cw, 3.9, fill=SURFACE, line=LINE)
    text(s, x + 0.22, 3.12, cw - 0.4, 0.3, k, size=11, color=ACCENT, font=MONO, bold=True)
    text(s, x + 0.22, 3.55, cw - 0.4, 3.2,
         [[("•  " + it, {"space_after": 9})] for it in items], size=13, color=TEXT, spacing=1.15)
notes(s, """
Under the hood this is a real product. The FastAPI backend owns the verdict. PostgreSQL stores claims, documents, pages, evidence and every verification with a snapshot of the evidence it used.
PDFs are parsed in a separate subprocess with a timeout and memory limit, so a hostile PDF cannot take the server down.
The AI layer is one interface with four providers behind it. Which one runs is configuration only, and each provider can have backup models that are tried in order.
Security basics are done properly: Argon2id passwords, hashed tokens, revocation, rate limits on authentication and OTP, and ownership checks so one user can never read another user's claims or documents.
""")

# ---------------------------------------------------------------- 8. Proof
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 8, "Proof it works", "Measured today, not estimated.")
stats = [
    ("297", "backend tests passing", "186 unit · 111 integration against real PostgreSQL"),
    ("47", "security tests", "cross-user access, token revocation, rate limits, body limits, lock and delete protection"),
    ("3 of 4", "AI providers verified live", "Gemini, NVIDIA NIM and Groq returned validated verdicts; Ollama not run here"),
]
for i, (big, label, sub) in enumerate(stats):
    x = 0.6 + i * 4.08
    box(s, x, 2.85, 3.9, 2.25, fill=SURFACE, line=LINE)
    text(s, x + 0.3, 3.0, 3.4, 0.9, big, size=46, color=ACCENT, font=HEAD, bold=True)
    text(s, x + 0.3, 3.92, 3.4, 0.4, label, size=15, font=HEAD, bold=True)
    text(s, x + 0.3, 4.3, 3.4, 0.8, sub, size=11.5, color=MUTED, spacing=1.15)
box(s, 0.6, 5.3, 12.1, 1.55, fill=SURFACE_2)
text(s, 0.9, 5.45, 11.5, 0.3, "END-TO-END RUN IN THE WEB APP  ·  27 SEP 2026", size=11, color=MUTED, font=MONO, bold=True)
flow = ["Register", "Email code", "Claim", "Upload PDF", "Pick passages", "Verify", "Verdict", "Show me why"]
x = 0.9
for j, f in enumerate(flow):
    w = pill(s, x, 5.95, f, TEXT, SURFACE, size=11)
    x += w + 0.08
    if j < len(flow) - 1:
        text(s, x, 5.95, 0.2, 0.34, "›", size=16, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        x += 0.22
text(s, 0.9, 6.45, 11.5, 0.3, "Every step completed against the real backend. Result: Contradicted, 97%, 3 passages cited by page.",
     size=12, color=MUTED)
notes(s, """
Here is what we actually verified today.
297 backend tests pass: 186 unit tests and 111 integration tests that run against a real PostgreSQL database. 47 of those are security tests.
Three AI providers returned validated verdicts in live calls today: Gemini, NVIDIA NIM and Groq. Ollama is implemented but we did not run it.
And we ran the full user journey in the browser, from registration and the email code to the verdict and the highlighted source.
Being honest about gaps: the web and mobile apps have no automated tests yet, and there are no usage or accuracy figures because the product has not been benchmarked.
""")

# ---------------------------------------------------------------- 9. Responsible AI
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 9, "Responsible AI", "The model reasons. The backend decides.")
controls = [
    ("Grounded in supplied evidence only", "The prompt forbids outside knowledge; the model only sees the passages you attached."),
    ("Prompt-injection guard", "Claim and evidence are wrapped as untrusted data; injected delimiter tags are stripped."),
    ("Structured, validated output", "Strict JSON schema, parsed with Pydantic; malformed output is retried, then rejected."),
    ("No invented citations", "References to unknown passages are dropped; quotes must appear verbatim in the stored passage."),
    ("Checks can overrule the model", "A number, date or name conflict turns an AI “supported” into contradicted or partial."),
    ("Honest failure", "Provider outage: deterministic result only if decisive, otherwise “unavailable”. No verdict is invented."),
]
for i, (t, d) in enumerate(controls):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.1, 2.85 + row * 1.35
    box(s, x, y, 5.95, 1.2, fill=SURFACE, line=LINE)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.28), Inches(y + 0.3), Inches(0.14), Inches(0.14))
    dot.fill.solid()
    dot.fill.fore_color.rgb = ACCENT
    dot.line.fill.background()
    text(s, x + 0.6, y + 0.2, 5.1, 0.35, t, size=15, font=HEAD, bold=True)
    text(s, x + 0.6, y + 0.58, 5.1, 0.6, d, size=12, color=MUTED, spacing=1.15)
notes(s, """
AI does real work in ProofLens, but it does not get the last word.
The model only sees the passages you supplied, and the prompt tells it to treat them as untrusted data, so a document cannot instruct the model.
Its answer must match a strict schema. We validate it. Citations that point at passages we never gave it are dropped, and quotes have to appear verbatim.
When the deterministic checks find a hard conflict, like the wrong number, they overrule a model that says supported.
And when the AI is down, we say so. We would rather show "unavailable" than a verdict nobody can stand behind.
Confidence is shown as how strongly the evidence settles the claim, not the probability the claim is true, and the app says exactly that on screen.
""")

# ---------------------------------------------------------------- 10. Next
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, 10, "Next step", "From a working product to one people rely on.")
box(s, 0.6, 2.85, 5.9, 3.3, fill=SURFACE, line=LINE)
text(s, 0.9, 3.05, 5.3, 0.3, "WORKING TODAY", size=11, color=GREEN, font=MONO, bold=True)
text(s, 0.9, 3.45, 5.4, 2.6, [[("•  " + t, {"space_after": 8})] for t in [
    "Claim → evidence → verdict → explanation → source, on web",
    "Four verdicts with deterministic checks + AI reasoning",
    "Page-level Show me why with highlighted passages",
    "Email-code accounts and verification history",
]], size=14, spacing=1.15)
box(s, 6.8, 2.85, 5.9, 3.3, fill=SURFACE, line=ACCENT, weight=1.5)
text(s, 7.1, 3.05, 5.3, 0.3, "NEXT  ·  NOT BUILT YET", size=11, color=ACCENT, font=MONO, bold=True)
text(s, 7.1, 3.45, 5.4, 2.6, [[("•  " + t, {"space_after": 8})] for t in [
    "Suggest the relevant passages automatically (ranking code exists, not yet wired in)",
    "Check that a passage really appears on the page it cites",
    "Split multi-part claims and verify each part",
    "Web and mobile test suites · accuracy benchmark",
]], size=14, spacing=1.15)
text(s, 0.6, 6.45, 12.1, 0.6, [
    [("ProofLens helps people say exactly what their evidence supports, ", {}), ("and show the page that proves it.", {"color": ACCENT})],
], size=20, font=HEAD, bold=True)
notes(s, """
Today ProofLens does the full loop: a claim, your evidence, a verdict with an explanation, and the exact source passage.
Next, in order: suggest the relevant passages automatically, since the page-ranking code already exists but is not yet connected; confirm that a chosen passage really appears on the page it cites; split claims with several parts and check each one; and add tests for the web and mobile apps plus a proper accuracy benchmark.
The goal stays the same: help people say exactly what their evidence supports, and show the page that proves it. Thank you.
""")

out = SP / "ProofLens-pitch.pptx"
prs.save(out)
print(out)
